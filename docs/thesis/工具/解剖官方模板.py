"""解剖川农官方专业型硕士论文模板（附录D）的结构、样式与页面设置。"""
import sys
from docx import Document

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

path = sys.argv[1]
d = Document(path)

print("=== 节（页面设置）===")
for i, s in enumerate(d.sections):
    print(f"[{i}] 页面 {s.page_width.cm:.1f}x{s.page_height.cm:.1f}cm | "
          f"边距 上{s.top_margin.cm:.2f} 下{s.bottom_margin.cm:.2f} "
          f"左{s.left_margin.cm:.2f} 右{s.right_margin.cm:.2f} | "
          f"页眉{s.header_distance.cm:.2f} 页脚{s.footer_distance.cm:.2f}")
    print(f"     首页不同={s.different_first_page_header_footer} | "
          f"evenAndOdd={d.settings.element.find(W + 'evenAndOddHeaders') is not None}")
    try:
        h = [p.text for p in s.header.paragraphs if p.text.strip()]
        f = [p.text for p in s.footer.paragraphs if p.text.strip()]
        print(f"     页眉: {h[:2]}")
        print(f"     页脚: {f[:2]}")
    except Exception as e:
        print("     页眉页脚读取失败:", type(e).__name__)

print()
print("=== 段落样式使用情况 ===")
used = {}
for p in d.paragraphs:
    if p.text.strip():
        used.setdefault(p.style.name, []).append(p.text.strip()[:44])
for k, v in used.items():
    print(f"{k}  (x{len(v)})  例: {v[0]}")

print()
print("=== 文档段落序列（前 80 段）===")
n = 0
for p in d.paragraphs:
    t = p.text.strip()
    if t:
        n += 1
        print(f"{n:3d} [{p.style.name}] {t[:70]}")
        if n >= 80:
            break

print()
print("=== 已定义样式（段落类）===")
for st in d.styles:
    if st.type is not None and "PARAGRAPH" in str(st.type):
        try:
            fnt = st.font
            sz = fnt.size.pt if fnt.size else None
            nm = fnt.name
            b = fnt.bold
        except Exception:
            sz = nm = b = None
        print(f"  {st.name:<28} 字号={sz} 字体={nm} 粗体={b}")
