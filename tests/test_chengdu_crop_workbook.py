from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from processing.chengdu_crop_workbook import extract_target_crop_workbook


def _write_workbook(path: Path) -> Path:
    workbook = openpyxl.Workbook()
    biomass = workbook.active
    biomass.title = "Pidu biomass"
    biomass.append(["date", "planting", "sample", "fresh", None, None, None, "secondary", None, None, None, None, "yield", None, None])
    biomass.append([None, None, None, "root", "stem", "leaf", "fruit", "root", "stem", "leaf", "fruit", None, "truss1", "truss2", "truss3"])
    for sample in (1, 2, 3):
        biomass.append([None, None, sample, 1.0, 2.0, 3.0, 0.0, 0.1, 0.2, 0.3, 0.0, None, 100.0, 120.0, 80.0])
    for date, fruit in (("2026-03-30", 0.0), ("2026-04-27", 250.0)):
        for sample in (1, 2, 3):
            biomass.append([date if sample == 1 else None, None, sample, 10.0, 20.0, 30.0, fruit + sample, 1.0, 2.0, 3.0, 4.0])

    morphology = workbook.create_sheet("Pidu morphology")
    morphology.append(["date", "planting", "sample", "height", "diameter", "nodes", "flowers", "fruit_trusses"])
    morphology.append([None] * 8)
    for sample in (1, 2, 3):
        morphology.append([None, None, sample, 13.0, 2.0, 6, 0, 0])
    for date, height in (("2026-03-30", 40.0), ("2026-04-27", 120.0)):
        for sample in (1, 2, 3):
            morphology.append([date if sample == 1 else None, None, sample, height + sample, 10.0, 24, 7, 4])
    workbook.save(path)
    return path


def test_extract_target_workbook_keeps_only_dated_standing_crop_samples(tmp_path: Path):
    observations, audit = extract_target_crop_workbook(
        _write_workbook(tmp_path / "crop.xlsx"),
        biomass_sheet="Pidu biomass",
        morphology_sheet="Pidu morphology",
        greenhouse_id=63,
        greenhouse_code="GH-PIDU",
        planting_id=59,
        planting_code="P-2026",
        cultivar="Sainirui",
        plant_density_plants_m2=1.0,
    )

    assert len(observations) == 6
    assert observations["target_eligible"].all()
    assert observations["observation_date"].nunique() == 2
    assert str(observations["observation_date"].dt.tz) == "Asia/Shanghai"
    assert observations["ripe_fruit_fresh_g_per_plant"].tolist()[-3:] == [251.0, 252.0, 253.0]
    assert observations["ripe_fruit_fresh_kg_m2"].tolist()[-3:] == pytest.approx([0.251, 0.252, 0.253])
    assert observations["plant_height_cm"].tolist()[-3:] == [121.0, 122.0, 123.0]
    assert observations["stem_node_count"].tolist()[-3:] == [24.0, 24.0, 24.0]
    assert observations["ripe_fruit_dry_g_per_plant"].isna().all()
    assert observations["quality_flags"].str.contains("secondary_mass_header_ambiguous_not_used").all()
    assert audit["dated_observation_rows"] == 6
    assert audit["undated_sample_rows_excluded"] == 3
    assert audit["undated_rows_with_reported_truss_yield"] == 3
    assert audit["observation_semantics"] == "standing_crop_sample_not_harvest_event"


def test_extract_target_workbook_rejects_negative_fresh_mass(tmp_path: Path):
    path = _write_workbook(tmp_path / "crop.xlsx")
    workbook = openpyxl.load_workbook(path)
    workbook["Pidu biomass"]["G9"] = -1.0
    workbook.save(path)

    with pytest.raises(ValueError, match="non-negative"):
        extract_target_crop_workbook(
            path,
            biomass_sheet="Pidu biomass",
            morphology_sheet="Pidu morphology",
            greenhouse_id=63,
            greenhouse_code="GH-PIDU",
            planting_id=59,
            planting_code="P-2026",
            cultivar="Sainirui",
            plant_density_plants_m2=1.0,
        )
